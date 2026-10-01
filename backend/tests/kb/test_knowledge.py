"""Knowledge base: chunking, hybrid retrieval quality, audience isolation, ingestion formats."""

import io

import pytest
import yaml
from sqlalchemy import text

from atrium.config import REPO_ROOT
from atrium.kb.chunking import MAX_CHARS, chunk_markdown, to_markdown
from atrium.kb.embeddings import HashEmbedder

EVAL = yaml.safe_load((REPO_ROOT / "shared/eval/retrieval.yaml").read_text())["queries"]


def test_chunks_follow_headings_and_overlap_long_sections():
    long_para = " ".join(["Regra importante sobre férias e saldo."] * 40)
    md = f"# Política\n\n*Documento fictício para demonstração.*\n\n## Fracionamento\n\n{long_para}\n\n{long_para}\n\n### Exceções\n\nTexto curto."
    chunks = chunk_markdown(md)
    assert {c.title for c in chunks} == {"Política"}
    assert chunks[0].heading == "Fracionamento" and chunks[-1].heading == "Fracionamento › Exceções"
    assert all(len(c.content) <= MAX_CHARS * 1.5 for c in chunks)
    assert len([c for c in chunks if c.heading == "Fracionamento"]) >= 2
    assert all("Documento fictício" not in c.content for c in chunks)


def test_hash_embedder_is_deterministic_and_normalized():
    a, b = HashEmbedder().embed(["férias em dobro", "férias em dobro"])
    assert a == b and abs(sum(x * x for x in a) - 1) < 1e-9


@pytest.mark.db
def test_retrieval_hit_at_3(services, identity, capsys):
    hits, misses = 0, []
    personas = {"lideranca": "gestora", "people-analytics": "hrbp"}
    for item in EVAL:
        who = next((personas[k] for k in item["kb"] if k in personas), "colaborador")
        results = services.kb.search(identity(who), item["q"], item["kb"], limit=3)
        if any(r.source.endswith(item["doc"]) for r in results):
            hits += 1
        else:
            misses.append((item["q"], item["doc"], [r.source for r in results]))
    rate = hits / len(EVAL)
    with capsys.disabled():
        print(f"\nretrieval hit@3 (hash embeddings + full-text): {hits}/{len(EVAL)} = {rate:.0%}")
        for m in misses:
            print("  miss:", m)
    assert rate >= 0.8


@pytest.mark.db
def test_restricted_knowledge_bases_do_not_exist_for_other_audiences(services, identity, db):
    q = "Qual o prefixo dos modelos de staging no dbt?"
    assert services.kb.search(identity("colaborador"), q, ["plataforma-dados"])  # Plataforma de Dados unit
    assert services.kb.search(identity("novata"), q, ["plataforma-dados"]) == []  # other unit: RLS hides the chunks
    assert services.kb.search(identity("colaborador"), "aprovar férias do time", ["lideranca"]) == []
    assert services.kb.search(identity("gestora"), "aprovar férias do time", ["lideranca"])
    with db.scoped(identity("novata").employee_id) as c:
        assert c.execute(text("SELECT count(*) FROM app.kb_chunks WHERE kb_id IN ('plataforma-dados', 'lideranca', 'people-analytics')")).scalar_one() == 0


def test_docx_and_pdf_uploads_become_headed_markdown():
    import docx
    from weasyprint import HTML

    d = docx.Document()
    d.add_heading("Guia de Plantão", level=1)
    d.add_heading("Escalonamento", level=2)
    d.add_paragraph("Acione o líder técnico em até 15 minutos.")
    buf = io.BytesIO()
    d.save(buf)
    md = to_markdown(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "guia.docx")
    chunks = chunk_markdown(md)
    assert chunks[0].title == "Guia de Plantão" and chunks[0].heading == "Escalonamento"

    pdf = HTML(string="<h1>Política de Viagem</h1><p>Diária máxima de hotel: R$ 650,00.</p>").write_pdf()
    assert "Diária máxima de hotel" in to_markdown(pdf, "application/pdf", "viagem.pdf")
